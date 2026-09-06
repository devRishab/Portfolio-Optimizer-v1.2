import streamlit as st
import yfinance as yf
import numpy as np
import pandas as pd
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import scipy.optimize as sco
from datetime import date
import io
import warnings
warnings.filterwarnings("ignore")

# ─────────────────────────────────────────────
#  PAGE CONFIG
# ─────────────────────────────────────────────
st.set_page_config(
    page_title="Portfolio Optimizer",
    page_icon="📈",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ─────────────────────────────────────────────
#  CSS
# ─────────────────────────────────────────────
st.markdown("""
<style>
#MainMenu                                 { visibility: hidden; }
header[data-testid="stHeader"]            { background: transparent !important;
                                            height: 0px !important;
                                            min-height: 0 !important; }
[data-testid="stToolbar"]                 { display: none !important; }
.stDeployButton                           { display: none !important; }
footer                                    { display: none !important; }
[data-testid="stSidebarCollapsedControl"] { display: flex !important;
                                            visibility: visible !important; }
.stApp                        { background: #f1f5f9; }
.block-container              { padding-top: 1.2rem !important;
                                padding-bottom: 1rem !important; }
[data-testid="stSidebar"]            { background: #1e293b !important; }
[data-testid="stSidebar"] label      { color: #cbd5e1 !important; }
[data-testid="stSidebar"] .stCaption { color: #94a3b8 !important; }
[data-testid="stSidebar"] p          { color: #cbd5e1 !important; }
[data-testid="stSidebar"] small      { color: #94a3b8 !important; }
[data-testid="stSidebar"] h1,
[data-testid="stSidebar"] h2,
[data-testid="stSidebar"] h3         { color: #f1f5f9 !important; }
[data-testid="stSidebar"] .stTextArea textarea {
    background: #0f172a !important; color: #e2e8f0 !important;
    border: 1px solid #475569 !important; border-radius: 8px !important; }
[data-testid="stSidebar"] .stTextInput input {
    background: #0f172a !important; color: #e2e8f0 !important;
    border: 1px solid #475569 !important; border-radius: 8px !important; }
[data-testid="stSidebar"] .stDateInput input {
    background: #0f172a !important; color: #e2e8f0 !important;
    border: 1px solid #475569 !important; }
[data-testid="stSidebar"] [data-testid="stSlider"] div { color: #cbd5e1 !important; }
[data-testid="stSidebar"] .stSelectbox div { color: #e2e8f0 !important; }

[data-testid="stMetric"]      { background: #ffffff; border: 1px solid #e2e8f0;
                                 border-radius: 12px; padding: 14px !important;
                                 box-shadow: 0 1px 3px rgba(0,0,0,.07); }
[data-testid="stMetricLabel"] { color: #64748b !important; font-size: 12px !important;
                                 font-weight: 600 !important; }
[data-testid="stMetricValue"] { color: #0f172a !important; font-size: 22px !important;
                                 font-weight: 800 !important; }
[data-testid="stMetricDelta"] { font-size: 12px !important; }

.stTabs [data-baseweb="tab-list"] { background: #ffffff; border-radius: 10px;
                                     padding: 4px; gap: 3px;
                                     border: 1px solid #e2e8f0;
                                     box-shadow: 0 1px 3px rgba(0,0,0,.06); }
.stTabs [data-baseweb="tab"]      { border-radius: 7px; color: #64748b !important;
                                     font-weight: 600; font-size: 13px;
                                     padding: 7px 14px; border: none !important; }
.stTabs [aria-selected="true"]    { background: #2563eb !important;
                                     color: #ffffff !important; }
h1,h2,h3,h4 { color: #0f172a !important; }
p           { color: #334155; }
[data-testid="stDataFrame"] { border-radius: 10px; border: 1px solid #e2e8f0;
                               box-shadow: 0 1px 3px rgba(0,0,0,.05); }
[data-testid="stExpander"]  { background: #ffffff; border: 1px solid #e2e8f0;
                               border-radius: 10px; }
hr { border-color: #e2e8f0 !important; margin: 0.8rem 0 !important; }
::-webkit-scrollbar        { width:5px; height:5px; }
::-webkit-scrollbar-track  { background:#f1f5f9; }
::-webkit-scrollbar-thumb  { background:#cbd5e1; border-radius:3px; }
.stButton>button { border-radius: 9px !important; font-weight: 700 !important;
                   font-size: 14px !important; height: 44px; }
</style>
""", unsafe_allow_html=True)

# ─────────────────────────────────────────────
#  THEME TOKENS
# ─────────────────────────────────────────────
C = {
    "bg"    : "#ffffff", "page"  : "#f1f5f9",
    "border": "#e2e8f0", "head"  : "#0f172a",
    "body"  : "#334155", "muted" : "#64748b",
    "blue"  : "#2563eb", "green" : "#16a34a",
    "red"   : "#dc2626", "amber" : "#d97706",
    "purple": "#7c3aed",
}
PALETTE = ["#2563eb","#16a34a","#d97706","#dc2626",
           "#7c3aed","#0891b2","#ea580c","#65a30d"]

# ─────────────────────────────────────────────
#  MARKET INDEX MAP  — friendly name → Yahoo ticker
# ─────────────────────────────────────────────
MARKET_MAP = {
    "Nifty 50"          : "^NSEI",
    "Nifty 500"         : "^CRSLDX",
    "Nifty 100"         : "^CNX100",
    "Nifty Bank"        : "^NSEBANK",
    "Nifty IT"          : "^CNXIT",
    "Nifty Next 50"     : "^NSMIDCP",
    "Nifty Pharma"      : "^CNXPHARMA",
    "Nifty Midcap 100"  : "^NSEMDCP100",
    "BSE Sensex"        : "^BSESN",
    "S&P 500 (US)"      : "^GSPC",
    "Nasdaq 100 (US)"   : "^NDX",
    "Dow Jones (US)"    : "^DJI",
}

# ─────────────────────────────────────────────
#  CHART HELPERS
# ─────────────────────────────────────────────
def apply_white_theme(fig, height=500, margin=None, hovermode="closest"):
    m = margin or dict(t=30, b=30, l=10, r=10)
    fig.update_layout(
        paper_bgcolor=C["bg"], plot_bgcolor=C["bg"],
        height=height, margin=m, hovermode=hovermode,
        font=dict(color=C["body"], family="Inter, sans-serif"),
        legend=dict(bgcolor=C["bg"], bordercolor=C["border"],
                    borderwidth=1, font=dict(size=11, color=C["body"]))
    )

def style_axes(fig, xkw=None, ykw=None):
    xd = dict(gridcolor=C["border"], zeroline=False,
               tickfont=dict(size=11, color=C["muted"]),
               linecolor=C["border"], showline=True,
               title_font=dict(color=C["body"], size=12))
    yd = dict(gridcolor=C["border"], zeroline=False,
               tickfont=dict(size=11, color=C["muted"]),
               linecolor=C["border"], showline=True,
               title_font=dict(color=C["body"], size=12))
    if xkw: xd.update(xkw)
    if ykw: yd.update(ykw)
    fig.update_xaxes(**xd)
    fig.update_yaxes(**yd)

# ─────────────────────────────────────────────
#  TICKER AUTO-RESOLVER
#  User types  TCS / INFY / AAPL  (no suffix needed)
#  We try:  TCS.NS  →  TCS.BO  →  TCS  (US / crypto)
# ─────────────────────────────────────────────
@st.cache_data(show_spinner=False, ttl=3600)
def resolve_ticker(raw: str) -> str | None:
    """Return the first valid Yahoo Finance ticker for a raw symbol."""
    raw = raw.strip().upper()
    # If user already added a suffix, honour it directly
    if "." in raw:
        try:
            d = yf.download(raw, period="5d", progress=False, auto_adjust=True)
            if not d.empty:
                return raw
        except Exception:
            pass
        return None

    # Otherwise try suffixes in order
    for suffix in [".NS", ".BO", ""]:
        candidate = raw + suffix
        try:
            d = yf.download(candidate, period="5d", progress=False, auto_adjust=True)
            if not d.empty:
                return candidate
        except Exception:
            continue
    return None


def short(t: str) -> str:
    """Strip exchange suffix for display."""
    return t.replace(".NS","").replace(".BO","").replace(".BSE","")

# ─────────────────────────────────────────────
#  FINANCE FUNCTIONS
# ─────────────────────────────────────────────
@st.cache_data(show_spinner=False)
def fetch_data(resolved_tickers, mkt, start, end):
    all_t = list(resolved_tickers) + [mkt]
    raw   = yf.download(all_t, start=start, end=end,
                        auto_adjust=True, progress=False)["Close"]
    raw.dropna(how="all", inplace=True)
    avail = [t for t in resolved_tickers if t in raw.columns]
    return raw[avail].dropna(), raw[mkt].dropna(), avail


def compute_stats(sd, md):
    ret, mkt = sd.pct_change().dropna(), md.pct_change().dropna()
    ret, mkt = ret.align(mkt, join="inner", axis=0)
    return (ret, mkt,
            ret.mean()*252, ret.cov()*252, ret.corr(),
            float(mkt.mean()*252), float(np.var(mkt)))


def port_perf(w, mu, sig, rf):
    r  = float(np.dot(w, mu))
    v  = float(np.sqrt(np.dot(w.T, np.dot(sig, w))))
    sh = (r - rf) / v if v > 0 else 0.0
    return r, v, sh


def optimize_portfolio(mu, sig, rf):
    n = len(mu)
    res = sco.minimize(
        lambda w: -port_perf(w, mu, sig, rf)[2],
        n*[1/n], method="SLSQP",
        bounds=[(0,1)]*n,
        constraints={"type":"eq","fun":lambda x: np.sum(x)-1}
    )
    return res.x


def find_gmvp(mu, sig):
    n = len(mu)
    res = sco.minimize(
        lambda w: float(np.sqrt(np.dot(w.T, np.dot(sig, w)))),
        n*[1/n], method="SLSQP",
        bounds=[(0,1)]*n,
        constraints={"type":"eq","fun":lambda x: np.sum(x)-1}
    )
    if res.success:
        return float(np.dot(res.x, mu)), float(res.fun)
    return None, None


def min_var_vol(mu, sig, target):
    n = len(mu)
    res = sco.minimize(
        lambda w: float(np.sqrt(np.dot(w.T, np.dot(sig, w)))),
        n*[1/n], method="SLSQP",
        bounds=[(0,1)]*n,
        constraints=[
            {"type":"eq","fun":lambda x: np.sum(x)-1},
            {"type":"eq","fun":lambda x, t=target: np.dot(x,mu)-t}
        ]
    )
    return float(res.fun) if res.success else None


def get_betas(returns, mkt_ret, mkt_var):
    return {c: float(np.cov(returns[c], mkt_ret)[0,1]/mkt_var)
            for c in returns.columns}


def get_port_beta(port_daily, mkt_ret):
    return float(np.cov(port_daily, mkt_ret)[0,1]/np.var(mkt_ret))


# ─────────────────────────────────────────────
#  EXCEL BUILDER
# ─────────────────────────────────────────────
def build_excel(stock_data, market_data, returns, mkt_ret,
                mean_ret, cov_mat, corr_mat, mkt_annual,
                available, market_choice, rf_rate,
                opt_w, opt_ret, opt_vol, opt_sh,
                port_beta, capm_ret, alpha_val, betas,
                gmvp_ret, gmvp_vol, start_date, end_date):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    wb  = openpyxl.Workbook()
    lbl = [short(t) for t in available]
    n   = len(available)

    # ── styles ────────────────────────────────────────────────────────
    def _fill(hex_): return PatternFill("solid", fgColor=hex_)
    def _font(hex_, bold=False, sz=10): return Font(bold=bold, color=hex_, size=sz)
    thin  = Side(style="thin", color="CBD5E0")
    BRD   = Border(left=thin, right=thin, top=thin, bottom=thin)
    CTR   = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LFT   = Alignment(horizontal="left",   vertical="center")
    RGT   = Alignment(horizontal="right",  vertical="center")

    F_HDR  = _fill("1E3A5F"); FN_HDR  = _font("FFFFFF", True)
    F_SUB  = _fill("2D4A7A"); FN_SUB  = _font("FFFFFF", True)
    F_STAT = _fill("EFF6FF"); FN_STAT = _font("1E3A5F", True)
    F_TTL  = _fill("0F172A"); FN_TTL  = _font("FFFFFF", True, 12)
    FN_NRM = _font("1E293B")
    FN_GRN = _font("16A34A", True)
    FN_RED = _font("DC2626", True)

    def hdr(ws, r, c, v, fill=F_HDR, fn=FN_HDR, aln=CTR):
        cl = ws.cell(row=r, column=c, value=v)
        cl.fill = fill; cl.font = fn; cl.alignment = aln; cl.border = BRD
        return cl

    def put(ws, r, c, v, fmt=None, fn=FN_NRM, fill=None, aln=RGT):
        cl = ws.cell(row=r, column=c, value=v)
        cl.font = fn; cl.alignment = aln; cl.border = BRD
        if fmt:  cl.number_format = fmt
        if fill: cl.fill = fill
        return cl

    def cw(ws, col, w):
        ws.column_dimensions[get_column_letter(col)].width = w

    # ═══════════════════════════════════════════════════════════
    #  SHEET 1 — Daily Data & Statistics
    # ═══════════════════════════════════════════════════════════
    ws1 = wb.active
    ws1.title = "Data, Covariance & Correlation"
    ws1.freeze_panes = "B4"

    # column indices
    PC        = 2           # price col start (B)
    MKT_PC    = PC + n      # market price col
    RC        = MKT_PC + 2  # return col start
    MKT_RC    = RC + n      # market return col
    SC        = MKT_RC + 3  # stats label col

    # title
    last_col = SC + n + 1
    ws1.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    c_ = ws1.cell(row=1, column=1,
                  value=f"Portfolio Data & Statistics  |  {start_date} to {end_date}  |  Market: {market_choice}")
    c_.font = FN_TTL; c_.fill = F_TTL; c_.alignment = CTR

    # row 2 headers
    hdr(ws1, 2, 1, "Date")
    for i, lb in enumerate(lbl):
        hdr(ws1, 2, PC+i, f"{lb}\nClose Price")
    hdr(ws1, 2, MKT_PC, f"{market_choice}\nClose Price")
    hdr(ws1, 2, RC-1, "Daily Returns", fill=F_SUB, fn=FN_SUB)
    for i, lb in enumerate(lbl):
        hdr(ws1, 2, RC+i, f"{lb}\nReturn")
    hdr(ws1, 2, MKT_RC, f"{market_choice}\nReturn")
    hdr(ws1, 2, SC,   "Statistic",    fill=F_STAT, fn=FN_STAT)
    for i, lb in enumerate(lbl):
        hdr(ws1, 2, SC+1+i, lb,       fill=F_STAT, fn=FN_STAT)
    hdr(ws1, 2, SC+1+n, market_choice, fill=F_STAT, fn=FN_STAT)

    # row 3 ticker names
    hdr(ws1, 3, 1, "Date", fill=F_SUB, fn=FN_SUB)
    for i, t in enumerate(available):
        hdr(ws1, 3, PC+i, t,    fill=F_SUB, fn=FN_SUB)
    hdr(ws1, 3, MKT_PC, market_choice, fill=F_SUB, fn=FN_SUB)

    # column widths
    cw(ws1, 1, 13)
    for c__ in range(PC, MKT_RC+2):
        cw(ws1, c__, 14)
    cw(ws1, SC,     26)
    for c__ in range(SC+1, SC+n+3):
        cw(ws1, c__, 16)

    # ── daily data rows ──────────────────────────────────────────
    aligned_idx = returns.index
    price_df    = stock_data.reindex(aligned_idx)
    mkt_price   = market_data.reindex(aligned_idx)
    DATA_ROW    = 4   # first data row

    for ri, dt in enumerate(aligned_idx):
        row = DATA_ROW + ri
        # date
        dc = ws1.cell(row=row, column=1, value=dt.date())
        dc.number_format = "DD-MMM-YYYY"; dc.alignment = CTR
        dc.border = BRD; dc.font = FN_NRM
        # stock prices — raw values (basis for return formulas)
        for i, t in enumerate(available):
            v = float(price_df.loc[dt, t]) if t in price_df.columns else None
            put(ws1, row, PC+i, v, fmt="0.00")
        # market price
        v = float(mkt_price.loc[dt]) if dt in mkt_price.index else None
        put(ws1, row, MKT_PC, v, fmt="0.00")
        # stock returns — Excel formula: (price / prev_price) - 1
        for i in range(n):
            pcol = get_column_letter(PC+i)
            if ri == 0:
                v = float(returns.iloc[ri, i]) if ri < len(returns) else 0
                put(ws1, row, RC+i, v, fmt="0.00000000")
            else:
                put(ws1, row, RC+i,
                    f"=({pcol}{row}/{pcol}{row-1})-1",
                    fmt="0.00000000")
        # market return formula
        mcol = get_column_letter(MKT_PC)
        if ri == 0:
            v = float(mkt_ret.iloc[0]) if len(mkt_ret) > 0 else 0
            put(ws1, row, MKT_RC, v, fmt="0.00000000")
        else:
            put(ws1, row, MKT_RC,
                f"=({mcol}{row}/{mcol}{row-1})-1",
                fmt="0.00000000")

    D_END = DATA_ROW + len(aligned_idx) - 1  # last data row

    # ── statistics block — Excel formulas for AVERAGE/STDEV/VAR ──
    stat_defs = [
        ("Average Daily Return",   "AVERAGE", "0.00000000"),
        ("Annualised Return",       "ANNRET",  "0.00%"),
        ("Daily Std Dev",          "STDEV",   "0.00000000"),
        ("Annualised Std Dev",     "ANNSTD",  "0.00%"),
        ("Daily Variance",         "VAR",     "0.00000000"),
        ("Annualised Variance",    "ANNVAR",  "0.00000000"),
    ]
    sr = DATA_ROW  # stats start same row as data
    for s_lbl, stype, fmt in stat_defs:
        hdr(ws1, sr, SC, s_lbl, fill=F_STAT, fn=FN_STAT, aln=LFT)
        for i in range(n):
            rc_ = get_column_letter(RC+i)
            rng = f"{rc_}{DATA_ROW+1}:{rc_}{D_END}"
            if   stype == "AVERAGE": f_ = f"=AVERAGE({rng})"
            elif stype == "ANNRET":  f_ = f"=AVERAGE({rng})*252"
            elif stype == "STDEV":   f_ = f"=STDEV({rng})"
            elif stype == "ANNSTD":  f_ = f"=STDEV({rng})*SQRT(252)"
            elif stype == "VAR":     f_ = f"=VAR({rng})"
            elif stype == "ANNVAR":  f_ = f"=VAR({rng})*252"
            put(ws1, sr, SC+1+i, f_, fmt=fmt)
        mrc_ = get_column_letter(MKT_RC)
        mrng = f"{mrc_}{DATA_ROW+1}:{mrc_}{D_END}"
        if   stype == "AVERAGE": mf = f"=AVERAGE({mrng})"
        elif stype == "ANNRET":  mf = f"=AVERAGE({mrng})*252"
        elif stype == "STDEV":   mf = f"=STDEV({mrng})"
        elif stype == "ANNSTD":  mf = f"=STDEV({mrng})*SQRT(252)"
        elif stype == "VAR":     mf = f"=VAR({mrng})"
        elif stype == "ANNVAR":  mf = f"=VAR({mrng})*252"
        put(ws1, sr, SC+1+n, mf, fmt=fmt)
        sr += 1

    sr += 1
    # VAR row index for market (used for Beta = Cov/Var formula)
    VAR_MKT_ROW = DATA_ROW + 4   # 5th stat = Daily Variance

    # ── Covariance with Market ─────────────────────────────────
    hdr(ws1, sr, SC, "Covariance with Market (Daily)", fill=F_HDR)
    sr += 1
    hdr(ws1, sr, SC,   "Stock",                 fill=F_STAT, fn=FN_STAT, aln=LFT)
    hdr(ws1, sr, SC+1, "Cov(stock, Mkt) Daily", fill=F_STAT, fn=FN_STAT)
    hdr(ws1, sr, SC+2, "Beta = Cov / Var(Mkt)", fill=F_STAT, fn=FN_STAT)
    hdr(ws1, sr, SC+3, "CAPM = Rf + B*(Rm-Rf)", fill=F_STAT, fn=FN_STAT)
    sr += 1
    mrc_col = get_column_letter(MKT_RC)
    mrng2   = f"{mrc_col}{DATA_ROW+1}:{mrc_col}{D_END}"
    var_mkt_cell = f"{get_column_letter(SC+1+n)}{VAR_MKT_ROW}"  # Daily Var of market
    ann_ret_mkt  = f"{get_column_letter(SC+1+n)}{DATA_ROW+1}"   # Ann Return of market

    BETA_CELLS = {}  # store cell ref per stock for Sheet2 use
    for i, t in enumerate(available):
        rc_ = get_column_letter(RC+i)
        rng = f"{rc_}{DATA_ROW+1}:{rc_}{D_END}"
        cov_ref  = f"{get_column_letter(SC+1)}{sr+i}"
        beta_ref = f"{get_column_letter(SC+2)}{sr+i}"
        BETA_CELLS[t] = beta_ref
        put(ws1, sr+i, SC,   lbl[i],  aln=LFT)
        # Cov — use Python value (COVARIANCE.S can fail in some locales)
        put(ws1, sr+i, SC+1, float(np.cov(returns[t], mkt_ret)[0,1]),
            fmt="0.00000000")
        # Beta = Cov / Var(mkt) — simple division formula, no @ risk
        put(ws1, sr+i, SC+2,
            f"={cov_ref}/{var_mkt_cell}",
            fmt="0.0000")
        # CAPM — arithmetic formula referencing beta cell
        put(ws1, sr+i, SC+3,
            f"={rf_rate}+{beta_ref}*({ann_ret_mkt}-{rf_rate})",
            fmt="0.00%")
    sr += n + 1

    # ── Covariance Matrix — Python values (avoids COVARIANCE.S locale issues)
    hdr(ws1, sr, SC, "Covariance Matrix — Annualised (computed)", fill=F_HDR)
    sr += 1
    hdr(ws1, sr, SC, "", fill=F_STAT, fn=FN_STAT)
    for i, lb in enumerate(lbl):
        hdr(ws1, sr, SC+1+i, lb, fill=F_STAT, fn=FN_STAT)
    sr += 1
    for i, t in enumerate(available):
        hdr(ws1, sr, SC, lbl[i], fill=F_STAT, fn=FN_STAT, aln=LFT)
        for j, t2 in enumerate(available):
            put(ws1, sr, SC+1+j, float(cov_mat.loc[t, t2]), fmt="0.00000000")
        sr += 1

    sr += 1

    # ── Correlation Matrix — Python values ──────────────────────
    hdr(ws1, sr, SC, "Correlation Matrix (computed)", fill=F_HDR)
    sr += 1
    hdr(ws1, sr, SC, "", fill=F_STAT, fn=FN_STAT)
    for i, lb in enumerate(lbl):
        hdr(ws1, sr, SC+1+i, lb, fill=F_STAT, fn=FN_STAT)
    sr += 1
    for i, t in enumerate(available):
        hdr(ws1, sr, SC, lbl[i], fill=F_STAT, fn=FN_STAT, aln=LFT)
        for j, t2 in enumerate(available):
            put(ws1, sr, SC+1+j, float(corr_mat.loc[t, t2]), fmt="0.000000")
        sr += 1

    # ═══════════════════════════════════════════════════════════
    #  SHEET 2 — SML
    # ═══════════════════════════════════════════════════════════
    ws2 = wb.create_sheet("SML Analysis")
    for c__, w__ in zip(range(1,6), [28,14,20,20,18]):
        cw(ws2, c__, w__)

    ws2.merge_cells("A1:E1")
    c_ = ws2.cell(row=1, column=1, value="Security Market Line — CAPM Analysis")
    c_.font=FN_TTL; c_.fill=F_TTL; c_.alignment=CTR

    # anchor params in col G/H (no @ risk — these are single-cell refs)
    ws2.cell(row=2, column=7, value="Risk-Free Rate").font = FN_STAT
    ws2.cell(row=2, column=8, value=rf_rate).number_format = "0.00%"
    ws2.cell(row=3, column=7, value="Market Return").font  = FN_STAT
    ws2.cell(row=3, column=8, value=mkt_annual).number_format = "0.00%"
    RF2 = "H2"; RM2 = "H3"  # anchor cells

    for c__, h in enumerate(
            ["Stock", "Beta", "CAPM Required Return",
             "Actual Annual Return", "Verdict"], start=1):
        hdr(ws2, 5, c__, h)

    # risk-free row
    put(ws2, 6, 1, "Risk-Free Asset", aln=LFT)
    put(ws2, 6, 2, 0,           fmt="0.0000")
    put(ws2, 6, 3, f"={RF2}",   fmt="0.00%")
    put(ws2, 6, 4, f"={RF2}",   fmt="0.00%")
    put(ws2, 6, 5, "—",         aln=CTR)

    # market row
    put(ws2, 7, 1, market_choice, aln=LFT)
    put(ws2, 7, 2, 1.0,          fmt="0.0000")
    put(ws2, 7, 3, f"={RM2}",    fmt="0.00%")
    put(ws2, 7, 4, f"={RM2}",    fmt="0.00%")
    put(ws2, 7, 5, "Market Portfolio", aln=CTR)

    for idx, t in enumerate(available):
        row = 8 + idx
        b_   = betas[t]
        act  = float(mean_ret[t])
        capm_= rf_rate + b_*(mkt_annual - rf_rate)
        under = act > capm_
        put(ws2, row, 1, lbl[idx], aln=LFT)
        put(ws2, row, 2, b_,       fmt="0.000000")
        # CAPM formula — arithmetic, no multi-cell range, no @ risk
        put(ws2, row, 3, f"={RF2}+B{row}*({RM2}-{RF2})", fmt="0.00%")
        put(ws2, row, 4, act,      fmt="0.00%")
        # Verdict formula — IF is universally safe
        vc = ws2.cell(row=row, column=5,
                      value=f'=IF(D{row}>C{row},"Undervalued","Overvalued")')
        vc.alignment=CTR; vc.border=BRD
        vc.font = FN_GRN if under else FN_RED

    # ═══════════════════════════════════════════════════════════
    #  SHEET 3 — Optimal Portfolio + CAL
    # ═══════════════════════════════════════════════════════════
    ws3 = wb.create_sheet("Optimal Portfolio & CAL")
    for c__, w__ in zip(range(1,10), [28,13,13,13,13,13,13,13,13]):
        cw(ws3, c__, w__)
    for c__ in range(10, 14):
        cw(ws3, c__, 18)

    ws3.merge_cells("A1:H1")
    c_ = ws3.cell(row=1, column=1,
                  value="Optimal Risky Portfolio — SLSQP Max-Sharpe Optimization")
    c_.font=FN_TTL; c_.fill=F_TTL; c_.alignment=CTR

    # ── anchor params (single cells — no @ risk) ──────────────
    def anc(ws, r, lbl_, val, fmt_):
        ws.cell(row=r, column=11, value=lbl_).font = FN_STAT
        cl = ws.cell(row=r, column=12, value=val)
        cl.number_format = fmt_
    anc(ws3, 2, "Rf (Risk-Free)",   rf_rate,    "0.00%")
    anc(ws3, 3, "Rm (Market Ret)",  mkt_annual, "0.00%")
    anc(ws3, 4, "Opt. Return",      opt_ret,    "0.00%")
    anc(ws3, 5, "Opt. Volatility",  opt_vol,    "0.00%")
    anc(ws3, 6, "Opt. Beta",        port_beta,  "0.0000")
    ws3.cell(row=7, column=11, value="Sharpe Ratio").font = FN_STAT
    ws3.cell(row=7, column=12,
             value=f"=(L4-L2)/L5").number_format = "0.0000"
    RF3="L2"; RM3="L3"; OR3="L4"; OV3="L5"; OB3="L6"; OS3="L7"

    # ── weights table ─────────────────────────────────────────
    for c__, h in enumerate(
            ["Stock","Weight","Annual Return","Annual Std Dev",
             "Beta","CAPM Return","Jensen Alpha","Active?"], start=1):
        hdr(ws3, 2, c__, h)

    for idx, t in enumerate(available):
        row = 3 + idx
        w_  = float(opt_w[available.index(t)])
        ret_= float(mean_ret[t])
        std_= float(np.sqrt(np.var(returns[t], ddof=1)*252))
        b_  = betas[t]
        put(ws3, row, 1, lbl[idx], aln=LFT)
        put(ws3, row, 2, w_,       fmt="0.00%")
        put(ws3, row, 3, ret_,     fmt="0.00%")
        put(ws3, row, 4, std_,     fmt="0.00%")
        put(ws3, row, 5, b_,       fmt="0.0000")
        # CAPM formula — arithmetic only
        put(ws3, row, 6, f"={RF3}+E{row}*({RM3}-{RF3})", fmt="0.00%")
        # Alpha = actual - CAPM
        put(ws3, row, 7, f"=C{row}-F{row}", fmt="+0.00%;-0.00%")
        ac = ws3.cell(row=row, column=8,
                      value=f'=IF(B{row}>0.1%,"Yes","No")')
        ac.alignment=CTR; ac.border=BRD
        ac.font = FN_GRN if w_ > 0.001 else FN_NRM

    sr3 = 4 + n

    # ── portfolio summary (formula-linked to anchors) ──────────
    hdr(ws3, sr3, 1, "Portfolio Summary", fill=F_HDR)
    sr3 += 1
    summ = [
        ("Expected Annual Return", f"={OR3}",                "0.00%"),
        ("Annual Volatility",      f"={OV3}",                "0.00%"),
        ("Sharpe Ratio",           f"={OS3}",                "0.0000"),
        ("Portfolio Beta",         f"={OB3}",                "0.0000"),
        ("CAPM Expected Return",   f"={RF3}+{OB3}*({RM3}-{RF3})", "0.00%"),
        ("Jensen Alpha",           f"={OR3}-({RF3}+{OB3}*({RM3}-{RF3}))", "+0.00%;-0.00%"),
        ("GMVP Return",            gmvp_ret if gmvp_ret else "N/A",  "0.00%"),
        ("GMVP Volatility",        gmvp_vol if gmvp_vol else "N/A",  "0.00%"),
        ("Risk-Free Rate",         f"={RF3}", "0.00%"),
        ("Market Return",          f"={RM3}", "0.00%"),
        ("Market Risk Premium",    f"={RM3}-{RF3}", "0.00%"),
    ]
    for s_l, s_v, s_f in summ:
        hdr(ws3, sr3, 1, s_l, fill=F_STAT, fn=FN_STAT, aln=LFT)
        if isinstance(s_v, str) and s_v.startswith("="):
            put(ws3, sr3, 2, s_v, fmt=s_f)
        elif isinstance(s_v, float):
            put(ws3, sr3, 2, s_v, fmt=s_f)
        else:
            put(ws3, sr3, 2, str(s_v), aln=LFT)
        sr3 += 1

    sr3 += 1

    # ── CAL Simulation Table — all Excel formulas ──────────────
    ws3.merge_cells(start_row=sr3, start_column=1, end_row=sr3, end_column=7)
    c_ = ws3.cell(row=sr3, column=1,
                  value="Capital Allocation Line — 101 Combinations (formula-driven)")
    c_.font=FN_TTL; c_.fill=F_HDR; c_.alignment=CTR
    sr3 += 1

    for c__, h in enumerate(
            ["Risk-Free Wt","Risky Wt","Expected Return",
             "Volatility","Beta","CAPM Req. Return","Sharpe Ratio"], start=1):
        hdr(ws3, sr3, c__, h)
    sr3 += 1

    for i in range(101):
        row = sr3 + i
        put(ws3, row, 1, i/100.0,       fmt="0%")
        put(ws3, row, 2, f"=1-A{row}",  fmt="0%")
        # Return = wrf*Rf + wr*OptRet — simple arithmetic
        put(ws3, row, 3, f"=A{row}*{RF3}+B{row}*{OR3}", fmt="0.00%")
        # Vol = wr * OptVol
        put(ws3, row, 4, f"=B{row}*{OV3}",              fmt="0.00%")
        # Beta = wr * OptBeta
        put(ws3, row, 5, f"=B{row}*{OB3}",              fmt="0.0000")
        # CAPM
        put(ws3, row, 6, f"={RF3}+E{row}*({RM3}-{RF3})", fmt="0.00%")
        # Sharpe — IF to avoid div/0
        put(ws3, row, 7, f"=IF(D{row}>0,(C{row}-{RF3})/D{row},0)", fmt="0.0000")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ─────────────────────────────────────────────
#  SIDEBAR
# ─────────────────────────────────────────────
with st.sidebar:
    st.markdown("""
    <div style='text-align:center; padding:16px 0 18px 0;'>
        <div style='font-size:38px; line-height:1;'>📊</div>
        <div style='font-size:18px; font-weight:800; color:#f8fafc; margin-top:8px;'>
            Portfolio Optimizer</div>
        <div style='font-size:11px; color:#94a3b8; margin-top:4px;'>
            SLSQP · CAPM · MPT</div>
    </div>
    <hr style='border-color:#334155; margin:0 0 14px 0;'>
    """, unsafe_allow_html=True)

    # ── Date Range ──
    st.markdown("<p style='color:#93c5fd; font-size:10px; font-weight:800; letter-spacing:1.5px; text-transform:uppercase; margin-bottom:6px;'>📅 DATE RANGE</p>", unsafe_allow_html=True)
    c1, c2 = st.columns(2)
    with c1:
        start_date = st.date_input("S", value=date(2020,1,1),
                                   max_value=date.today(), label_visibility="collapsed")
        st.caption("Start date")
    with c2:
        end_date = st.date_input("E", value=date.today(),
                                 max_value=date.today(), label_visibility="collapsed")
        st.caption("End date")

    st.markdown("<hr style='border-color:#334155; margin:10px 0;'>", unsafe_allow_html=True)

    # ── Stock Tickers ──
    st.markdown("<p style='color:#93c5fd; font-size:10px; font-weight:800; letter-spacing:1.5px; text-transform:uppercase; margin-bottom:4px;'>📈 STOCK TICKERS</p>", unsafe_allow_html=True)
    st.caption("Just type the symbol — TCS, INFY, AAPL. One per line. No .NS needed.")

    ticker_input = st.text_area("Tickers", label_visibility="collapsed", height=210,
        value="TCS\nINFY\nHDFCBANK\nICICIBANK\nRELIANCE\n"
              "ITC\nSUNPHARMA\nBHARTIARTL\nM&M\nGOLDBEES")
    raw_tickers = [t.strip().upper() for t in ticker_input.strip().split("\n") if t.strip()]

    st.markdown("<hr style='border-color:#334155; margin:10px 0;'>", unsafe_allow_html=True)

    # ── Market Index ── friendly dropdown
    st.markdown("<p style='color:#93c5fd; font-size:10px; font-weight:800; letter-spacing:1.5px; text-transform:uppercase; margin-bottom:4px;'>🏦 MARKET INDEX</p>", unsafe_allow_html=True)

    market_choice = st.selectbox(
        "Market Index",
        options=list(MARKET_MAP.keys()),
        index=0,          # default: Nifty 50
        label_visibility="collapsed"
    )
    market_ticker = MARKET_MAP[market_choice]
    st.caption(f"Yahoo Finance ticker: `{market_ticker}`")

    st.markdown("<div style='height:4px'></div>", unsafe_allow_html=True)

    # ── Risk-Free Rate ──
    st.markdown("<p style='color:#93c5fd; font-size:10px; font-weight:800; letter-spacing:1.5px; text-transform:uppercase; margin-bottom:4px;'>💰 RISK-FREE RATE</p>", unsafe_allow_html=True)
    rf_rate = st.slider("Risk-Free Rate (%)", 0.0, 15.0, 6.5, 0.1) / 100

    st.markdown("<div style='height:8px'></div>", unsafe_allow_html=True)
    run_btn = st.button("🚀  Run Optimization", use_container_width=True, type="primary")

    st.markdown("<hr style='border-color:#334155; margin:10px 0;'>", unsafe_allow_html=True)
    st.markdown("<p style='color:#93c5fd; font-size:10px; font-weight:800; letter-spacing:1.5px; text-transform:uppercase; margin-bottom:4px;'>📥 EXCEL EXPORT</p>", unsafe_allow_html=True)
    want_excel = st.checkbox("Generate Excel report", value=False)

    st.markdown("""
    <div style='margin-top:12px; background:#0f172a; border:1px solid #334155;
                border-radius:10px; padding:10px 12px; font-size:11px; color:#94a3b8;'>
        <b style='color:#60a5fa;'>💡 Tips:</b><br>
        • Just enter symbols: <b>TCS</b>, <b>AAPL</b><br>
        • NSE stocks auto-detected (.NS)<br>
        • US stocks also supported<br>
        • 3–20 stocks work best
    </div>
    """, unsafe_allow_html=True)


# ─────────────────────────────────────────────
#  MAIN HEADER
# ─────────────────────────────────────────────
st.markdown("""
<div style='background:#ffffff; border:1px solid #e2e8f0; border-radius:14px;
            padding:16px 22px; margin-bottom:18px;
            box-shadow:0 1px 4px rgba(0,0,0,.06);'>
    <div style='display:flex; align-items:center; gap:12px;'>
        <div style='font-size:28px;'>📈</div>
        <div>
            <div style='font-size:22px; font-weight:800; color:#0f172a; line-height:1.2;'>
                Portfolio Optimization Dashboard</div>
            <div style='font-size:12px; color:#64748b; margin-top:3px;'>
                Modern Portfolio Theory &nbsp;·&nbsp; CAPM &nbsp;·&nbsp;
                Efficient Frontier &nbsp;·&nbsp; SLSQP Optimizer</div>
        </div>
    </div>
</div>
""", unsafe_allow_html=True)

# ─────────────────────────────────────────────
#  LANDING
# ─────────────────────────────────────────────
if not run_btn:
    st.markdown("""
    <div style='background:#ffffff; border:1px solid #e2e8f0; border-radius:16px;
                padding:48px 36px; text-align:center;
                box-shadow:0 1px 4px rgba(0,0,0,.06);'>
        <div style='font-size:56px; margin-bottom:12px;'>🏦</div>
        <div style='font-size:20px; font-weight:800; color:#0f172a; margin-bottom:8px;'>
            Configure your portfolio on the left panel</div>
        <div style='color:#64748b; font-size:13px; max-width:480px; margin:0 auto 24px;'>
            Type ticker symbols (no suffix needed) · Select market index ·
            Set date range · Click <b style='color:#2563eb;'>Run Optimization</b>
        </div>
        <div style='display:flex; justify-content:center; gap:10px; flex-wrap:wrap;'>
            <span style='background:#eff6ff; color:#1d4ed8; padding:6px 14px;
                         border-radius:999px; font-size:12px; font-weight:700;'>
                📊 Efficient Frontier</span>
            <span style='background:#f0fdf4; color:#15803d; padding:6px 14px;
                         border-radius:999px; font-size:12px; font-weight:700;'>
                📉 Capital Allocation Line</span>
            <span style='background:#faf5ff; color:#7e22ce; padding:6px 14px;
                         border-radius:999px; font-size:12px; font-weight:700;'>
                🎯 Security Market Line</span>
            <span style='background:#fff7ed; color:#c2410c; padding:6px 14px;
                         border-radius:999px; font-size:12px; font-weight:700;'>
                🔥 Correlation Matrix</span>
            <span style='background:#fefce8; color:#a16207; padding:6px 14px;
                         border-radius:999px; font-size:12px; font-weight:700;'>
                📋 CAPM Valuation</span>
        </div>
    </div>
    """, unsafe_allow_html=True)
    st.stop()


# ─────────────────────────────────────────────
#  RESOLVE TICKERS  (auto-detect exchange suffix)
# ─────────────────────────────────────────────
with st.spinner("🔍 Resolving ticker symbols..."):
    resolve_progress = st.empty()
    resolved_map   = {}   # raw → resolved ticker (e.g. TCS → TCS.NS)
    failed_tickers = []

    for raw in raw_tickers:
        resolved = resolve_ticker(raw)
        if resolved:
            resolved_map[raw] = resolved
        else:
            failed_tickers.append(raw)

    tickers = list(resolved_map.values())   # resolved tickers for download

if failed_tickers:
    st.warning(f"⚠️ Could not find: **{', '.join(failed_tickers)}** — skipped. "
               f"Check spelling or try adding suffix manually (e.g. TCS.NS).")

if len(tickers) < 2:
    st.error("❌ Need at least 2 valid tickers. Please check your inputs.")
    st.stop()




# ─────────────────────────────────────────────
#  FETCH & COMPUTE
# ─────────────────────────────────────────────
with st.spinner(f"⏳ Fetching data from {start_date} to {end_date} and running SLSQP..."):
    try:
        stock_data, market_data, available = fetch_data(
            tuple(tickers), market_ticker, str(start_date), str(end_date))

        if len(available) < 2:
            st.error("❌ Not enough data returned. Try a different date range.")
            st.stop()

        (returns, mkt_ret, mean_ret,
         cov_mat, corr_mat, mkt_annual, mkt_var) = compute_stats(stock_data, market_data)

        opt_w                    = optimize_portfolio(mean_ret, cov_mat, rf_rate)
        opt_ret, opt_vol, opt_sh = port_perf(opt_w, mean_ret, cov_mat, rf_rate)
        gmvp_ret, gmvp_vol       = find_gmvp(mean_ret, cov_mat)

        port_daily = returns.dot(opt_w)
        port_beta  = get_port_beta(port_daily, mkt_ret)
        capm_ret   = rf_rate + port_beta*(mkt_annual - rf_rate)
        alpha      = opt_ret - capm_ret
        betas      = get_betas(returns, mkt_ret, mkt_var)

    except Exception as e:
        st.error(f"❌ Error: {e}")
        st.stop()


# ─────────────────────────────────────────────
#  KEY METRICS
# ─────────────────────────────────────────────
st.markdown("<p style='font-size:15px; font-weight:700; color:#0f172a; margin-bottom:10px;'>🎯 Optimal Portfolio — Key Metrics</p>", unsafe_allow_html=True)

m1,m2,m3,m4,m5,m6 = st.columns(6)
active_n = sum(1 for w in opt_w if w > 0.001)
m1.metric("📈 Annual Return",  f"{opt_ret:.2%}",   delta=f"+{opt_ret-rf_rate:.2%} vs Rf")
m2.metric("📉 Volatility",     f"{opt_vol:.2%}",   delta="Total Risk")
m3.metric("⚡ Sharpe Ratio",   f"{opt_sh:.4f}",    delta="Max Sharpe")
m4.metric("🔵 Beta",           f"{port_beta:.4f}", delta="Systematic Risk")
m5.metric("🏦 CAPM Return",    f"{capm_ret:.2%}",  delta=f"Alpha: {alpha:+.2%}")
m6.metric("✅ Active Stocks",  f"{active_n}/{len(available)}", delta="Non-zero weights")

st.markdown("<div style='height:4px'></div>", unsafe_allow_html=True)
st.markdown("---")


# ─────────────────────────────────────────────
#  TABS
# ─────────────────────────────────────────────
_tab_labels = [
    "🏆  Optimal Portfolio",
    "📊  CAL & Efficient Frontier",
    "🎯  Security Market Line",
    "🔥  Correlation Matrix",
    "📉  Stock Valuation",
    "📋  CAL Simulation",
]
if want_excel:
    _tab_labels.append("📥  Export to Excel")

_tabs = st.tabs(_tab_labels)
tab1, tab2, tab3, tab4, tab5, tab6 = _tabs[:6]
tab7 = _tabs[6] if want_excel else None


# ════════════════════════════════════════════
#  TAB 1 — OPTIMAL PORTFOLIO
# ════════════════════════════════════════════
with tab1:
    L, R = st.columns([1.1, 1], gap="large")

    with L:
        st.markdown("<p style='font-size:14px; font-weight:700; color:#0f172a;'>📦 Asset Allocation Weights</p>", unsafe_allow_html=True)

        wdf = pd.DataFrame({
            "Ticker": available,
            "Label":  [short(t) for t in available],
            "Weight": opt_w,
            "Return": [float(mean_ret[t]) for t in available],
            "Beta":   [betas[t] for t in available],
        }).sort_values("Weight", ascending=False).reset_index(drop=True)

        adf    = wdf[wdf["Weight"] > 0.001].copy()
        colors = PALETTE[:len(adf)]

        fig_bar = go.Figure(go.Bar(
            x=adf["Label"], y=adf["Weight"]*100,
            marker=dict(color=colors, line=dict(color="#fff", width=2)),
            text=[f"{w*100:.1f}%" for w in adf["Weight"]],
            textposition="outside",
            textfont=dict(color=C["head"], size=12),
            hovertemplate="<b>%{x}</b><br>Weight: %{y:.2f}%<extra></extra>"
        ))
        apply_white_theme(fig_bar, height=280, margin=dict(t=10,b=10,l=10,r=10))
        style_axes(fig_bar,
                   xkw=dict(tickfont=dict(size=12, color=C["body"])),
                   ykw=dict(title="Weight (%)", ticksuffix="%"))
        fig_bar.update_layout(showlegend=False)
        st.plotly_chart(fig_bar, use_container_width=True)

        tbl = pd.DataFrame({
            "Ticker":       [short(t) for t in wdf["Ticker"]],
            "Weight %":     [f"{w:.2%}" for w in wdf["Weight"]],
            "Exp. Return":  [f"{r:.2%}" for r in wdf["Return"]],
            "Beta":         [f"{b:.4f}" for b in wdf["Beta"]],
            "In Portfolio": ["✅ Yes" if w > 0.001 else "⭕ No" for w in wdf["Weight"]]
        })
        st.dataframe(tbl, use_container_width=True, hide_index=True, height=275)

    with R:
        st.markdown("<p style='font-size:14px; font-weight:700; color:#0f172a;'>🥧 Portfolio Composition</p>", unsafe_allow_html=True)

        fig_pie = go.Figure(go.Pie(
            labels=adf["Label"],
            values=(adf["Weight"]*100).round(2),
            hole=0.56,
            marker=dict(colors=colors, line=dict(color="#fff", width=3)),
            textinfo="label+percent",
            textfont=dict(size=11, color=C["head"]),
            insidetextorientation="horizontal",
            hovertemplate="<b>%{label}</b><br>Weight: %{value:.2f}%<extra></extra>",
            direction="clockwise", sort=True
        ))
        apply_white_theme(fig_pie, height=280, margin=dict(t=10,b=10,l=10,r=10))
        fig_pie.update_layout(
            showlegend=False,
            annotations=[dict(
                text=f"<b>{opt_sh:.2f}</b><br><span style='font-size:11px'>Sharpe</span>",
                x=0.5, y=0.5, showarrow=False,
                font=dict(size=18, color=C["head"])
            )]
        )
        st.plotly_chart(fig_pie, use_container_width=True)

        st.markdown(f"""
        <div style='background:#ffffff; border:1px solid #e2e8f0; border-radius:14px;
                    padding:16px 18px; box-shadow:0 1px 4px rgba(0,0,0,.06);'>
            <p style='color:#64748b; font-size:10px; font-weight:800;
                      letter-spacing:1.5px; margin:0 0 12px 0;'>PORTFOLIO SUMMARY</p>
            <div style='display:grid; grid-template-columns:1fr 1fr; gap:12px;'>
                <div><div style='color:#64748b; font-size:11px;'>Annual Return</div>
                     <div style='color:#16a34a; font-size:20px; font-weight:800;'>{opt_ret:.2%}</div></div>
                <div><div style='color:#64748b; font-size:11px;'>Volatility</div>
                     <div style='color:#d97706; font-size:20px; font-weight:800;'>{opt_vol:.2%}</div></div>
                <div><div style='color:#64748b; font-size:11px;'>Sharpe Ratio</div>
                     <div style='color:#2563eb; font-size:20px; font-weight:800;'>{opt_sh:.4f}</div></div>
                <div><div style='color:#64748b; font-size:11px;'>Beta</div>
                     <div style='color:#7c3aed; font-size:20px; font-weight:800;'>{port_beta:.4f}</div></div>
                <div><div style='color:#64748b; font-size:11px;'>CAPM Return</div>
                     <div style='color:#dc2626; font-size:20px; font-weight:800;'>{capm_ret:.2%}</div></div>
                <div><div style='color:#64748b; font-size:11px;'>Jensen's Alpha</div>
                     <div style='color:{"#16a34a" if alpha>=0 else "#dc2626"}; font-size:20px; font-weight:800;'>{alpha:+.2%}</div></div>
            </div>
        </div>
        """, unsafe_allow_html=True)


# ════════════════════════════════════════════
#  TAB 2 — CAL & EFFICIENT FRONTIER
# ════════════════════════════════════════════
with tab2:
    st.markdown("<p style='font-size:14px; font-weight:700; color:#0f172a; margin-bottom:2px;'>📊 Capital Allocation Line & Efficient Frontier</p>", unsafe_allow_html=True)
    st.caption("5,000 random portfolios · Efficient frontier (upper portion from GMVP) · CAL from Rf through tangency point")

    with st.spinner("Generating efficient frontier..."):
        np.random.seed(42)
        n = len(available)
        sim_r, sim_v, sim_s = [], [], []
        for _ in range(5000):
            w = np.random.random(n); w /= w.sum()
            r, v, s = port_perf(w, mean_ret, cov_mat, rf_rate)
            sim_r.append(r); sim_v.append(v); sim_s.append(s)
        sim_r = np.array(sim_r)
        sim_v = np.array(sim_v)
        sim_s = np.array(sim_s)

        ef_start = gmvp_ret if gmvp_ret is not None else float(mean_ret.min())
        ef_end   = float(mean_ret.max())
        ef_v, ef_r = [], []
        for tr in np.linspace(ef_start, ef_end, 60):
            v = min_var_vol(mean_ret, cov_mat, tr)
            if v is not None:
                ef_v.append(v); ef_r.append(tr)

    cal_max = opt_vol * 1.55
    cal_x   = np.linspace(0, cal_max, 100)
    cal_y   = rf_rate + opt_sh * cal_x

    fig_cal = go.Figure()

    fig_cal.add_trace(go.Scatter(
        x=sim_v*100, y=sim_r*100, mode="markers",
        marker=dict(
            color=sim_s,
            colorscale=[[0,"#bfdbfe"],[0.5,"#3b82f6"],[1,"#1e3a8a"]],
            size=3.5, opacity=0.5,
            colorbar=dict(
                title=dict(text="Sharpe", font=dict(color=C["body"], size=11)),
                tickfont=dict(color=C["body"], size=10),
                x=1.01, thickness=12, len=0.65,
                bgcolor=C["bg"], bordercolor=C["border"]
            ), showscale=True
        ),
        name="Random Portfolios",
        hovertemplate="Vol: %{x:.2f}%<br>Return: %{y:.2f}%<extra>Random Portfolio</extra>"
    ))

    if len(ef_v) > 3:
        fig_cal.add_trace(go.Scatter(
            x=[v*100 for v in ef_v], y=[r*100 for r in ef_r],
            mode="lines", line=dict(color=C["amber"], width=3.5),
            name="Efficient Frontier",
            hovertemplate="Vol: %{x:.2f}%<br>Return: %{y:.2f}%<extra>Efficient Frontier</extra>"
        ))

    if gmvp_ret and gmvp_vol:
        fig_cal.add_trace(go.Scatter(
            x=[gmvp_vol*100], y=[gmvp_ret*100], mode="markers+text",
            marker=dict(color=C["amber"], size=11, symbol="diamond",
                        line=dict(color="white", width=2)),
            text=["GMVP"], textposition="bottom right",
            textfont=dict(color=C["amber"], size=11),
            name=f"GMVP ({gmvp_ret:.1%})",
            hovertemplate=(f"<b>GMVP</b><br>Return: {gmvp_ret:.2%}<br>"
                           f"Vol: {gmvp_vol:.2%}<extra></extra>")
        ))

    fig_cal.add_trace(go.Scatter(
        x=cal_x*100, y=cal_y*100, mode="lines",
        line=dict(color=C["blue"], width=2.5, dash="dash"),
        name=f"CAL (Sharpe={opt_sh:.2f})",
        hovertemplate="Vol: %{x:.2f}%<br>Return: %{y:.2f}%<extra>CAL</extra>"
    ))

    fig_cal.add_trace(go.Scatter(
        x=[0], y=[rf_rate*100], mode="markers+text",
        marker=dict(color=C["blue"], size=10, symbol="circle",
                    line=dict(color="white", width=2)),
        text=[f"Rf={rf_rate:.1%}"], textposition="top right",
        textfont=dict(color=C["blue"], size=11),
        name=f"Risk-Free ({rf_rate:.1%})",
        hovertemplate=f"Risk-Free: {rf_rate:.2%}<extra></extra>"
    ))

    fig_cal.add_trace(go.Scatter(
        x=[opt_vol*100], y=[opt_ret*100], mode="markers+text",
        marker=dict(color=C["red"], size=18, symbol="star",
                    line=dict(color="white", width=1.5)),
        text=["Optimal"], textposition="top right",
        textfont=dict(color=C["red"], size=12),
        name=f"Optimal (Sharpe={opt_sh:.2f})",
        hovertemplate=(f"<b>Optimal Portfolio</b><br>"
                       f"Return: {opt_ret:.2%}<br>Vol: {opt_vol:.2%}<br>"
                       f"Sharpe: {opt_sh:.4f}<extra></extra>")
    ))

    x_max = max(float(np.percentile(sim_v,99)), opt_vol)*100*1.08
    y_min = rf_rate*100*0.7
    y_max = max(float(np.percentile(sim_r,99)), opt_ret)*100*1.08

    apply_white_theme(fig_cal, height=550, margin=dict(t=20,b=20,l=10,r=70))
    style_axes(fig_cal,
               xkw=dict(title="Annual Volatility (%)", ticksuffix="%", range=[0, x_max]),
               ykw=dict(title="Expected Annual Return (%)", ticksuffix="%", range=[y_min, y_max]))
    fig_cal.update_layout(hovermode="closest")
    st.plotly_chart(fig_cal, use_container_width=True)

    with st.expander("📖 How to read this chart"):
        a, b = st.columns(2)
        with a:
            st.markdown("""
            - 🔵 **Blue dots** = 5,000 random portfolios. Darker = higher Sharpe.
            - 🟡 **Yellow curve** = Efficient Frontier — upper portion only (GMVP upward).
            - 🔷 **Yellow diamond** = GMVP — lowest-risk portfolio.
            """)
        with b:
            st.markdown("""
            - 🔵 **Blue dashed line** = CAL — Rf through optimal portfolio. Same Sharpe everywhere.
            - ⭐ **Red star** = Optimal (Tangency) Portfolio — maximum Sharpe.
            - 🔵 **Blue dot** = Risk-Free Rate at zero volatility.
            """)

    s1,s2,s3,s4 = st.columns(4)
    s1.metric("Optimal Return",     f"{opt_ret:.2%}")
    s2.metric("Optimal Volatility", f"{opt_vol:.2%}")
    s3.metric("Max Sharpe Ratio",   f"{opt_sh:.4f}")
    if gmvp_ret and gmvp_vol:
        s4.metric("GMVP", f"{gmvp_ret:.2%} ret | {gmvp_vol:.2%} vol")


# ════════════════════════════════════════════
#  TAB 3 — SECURITY MARKET LINE
# ════════════════════════════════════════════
with tab3:
    st.markdown("<p style='font-size:14px; font-weight:700; color:#0f172a; margin-bottom:2px;'>🎯 Security Market Line — Individual Stock Positioning</p>", unsafe_allow_html=True)

    beta_range  = np.linspace(0, max(betas.values())*1.25, 120)
    sml_returns = rf_rate + beta_range*(mkt_annual - rf_rate)

    fig_sml = go.Figure()
    fig_sml.add_trace(go.Scatter(
        x=beta_range, y=sml_returns*100, mode="lines",
        line=dict(color=C["blue"], width=2.5),
        name="Security Market Line",
        hovertemplate="Beta: %{x:.2f}<br>CAPM Return: %{y:.2f}%<extra>SML</extra>"
    ))

    for ticker in available:
        b      = betas[ticker]
        actual = float(mean_ret[ticker])
        capm_r = rf_rate + b*(mkt_annual - rf_rate)
        under  = actual > capm_r
        lbl    = short(ticker)
        fig_sml.add_trace(go.Scatter(
            x=[b], y=[actual*100], mode="markers+text",
            marker=dict(color=C["green"] if under else C["red"],
                        size=12, symbol="circle",
                        line=dict(color="white", width=1.5)),
            text=[lbl], textposition="top center",
            textfont=dict(size=10, color=C["head"]),
            name=lbl, showlegend=False,
            hovertemplate=(f"<b>{ticker}</b><br>Beta: {b:.4f}<br>"
                           f"Actual: {actual:.2%}<br>CAPM: {capm_r:.2%}<br>"
                           f"Alpha: {actual-capm_r:+.2%}<br>"
                           f"{'✅ Undervalued' if under else '❌ Overvalued'}<extra></extra>")
        ))

    fig_sml.add_trace(go.Scatter(
        x=[port_beta], y=[capm_ret*100], mode="markers+text",
        marker=dict(color=C["red"], size=17, symbol="star",
                    line=dict(color="white", width=1.5)),
        text=["Optimal"], textposition="top right",
        textfont=dict(color=C["red"], size=12),
        name="Optimal Portfolio",
        hovertemplate=f"Beta: {port_beta:.4f}<br>CAPM: {capm_ret:.2%}<extra>Optimal</extra>"
    ))
    fig_sml.add_trace(go.Scatter(
        x=[1.0], y=[mkt_annual*100], mode="markers+text",
        marker=dict(color=C["purple"], size=13, symbol="square",
                    line=dict(color="white", width=1.5)),
        text=["Market"], textposition="top right",
        textfont=dict(color=C["purple"], size=12),
        name=f"Market ({market_choice})",
        hovertemplate=f"Market Return: {mkt_annual:.2%}<extra>Market</extra>"
    ))
    fig_sml.add_trace(go.Scatter(
        x=[0], y=[rf_rate*100], mode="markers+text",
        marker=dict(color=C["blue"], size=10, symbol="circle",
                    line=dict(color="white", width=2)),
        text=["Rf"], textposition="top right",
        textfont=dict(color=C["blue"], size=12),
        name=f"Risk-Free ({rf_rate:.1%})",
        hovertemplate=f"Rf: {rf_rate:.2%}<extra></extra>"
    ))

    apply_white_theme(fig_sml, height=530)
    style_axes(fig_sml,
               xkw=dict(title="Beta (Systematic Risk)"),
               ykw=dict(title="Expected Return (%)", ticksuffix="%"))
    fig_sml.update_layout(hovermode="closest")
    st.plotly_chart(fig_sml, use_container_width=True)

    with st.expander("📖 How to read this chart"):
        a, b = st.columns(2)
        with a:
            st.markdown("""
            - **Blue line** = SML — CAPM required return for each beta level
            - 🟢 **Green dots** = Undervalued (actual > CAPM → positive alpha)
            - 🔴 **Red dots** = Overvalued (actual < CAPM → negative alpha)
            """)
        with b:
            st.markdown("""
            - 🟣 **Purple square** = Market portfolio (beta = 1 by definition)
            - ⭐ **Red star** = Optimal portfolio on SML
            - Hover over any dot for full alpha details
            """)


# ════════════════════════════════════════════
#  TAB 4 — CORRELATION MATRIX
# ════════════════════════════════════════════
with tab4:
    st.markdown("<p style='font-size:14px; font-weight:700; color:#0f172a; margin-bottom:2px;'>🔥 Correlation Matrix — Diversification Analysis</p>", unsafe_allow_html=True)

    slbls = [short(t) for t in available]
    cvals = corr_mat.values

    fig_heat = go.Figure(go.Heatmap(
        z=cvals, x=slbls, y=slbls,
        colorscale=[
            [0.00,"#1e40af"],[0.35,"#93c5fd"],
            [0.50,"#f8fafc"],[0.65,"#fca5a5"],[1.00,"#dc2626"]
        ],
        zmin=-1, zmax=1,
        text=np.round(cvals,2), texttemplate="%{text}",
        textfont=dict(size=11, color=C["head"]),
        hoverongaps=False,
        hovertemplate="<b>%{y} vs %{x}</b><br>Correlation: %{z:.4f}<extra></extra>",
        colorbar=dict(
            title=dict(text="Correlation", font=dict(color=C["body"], size=12)),
            tickfont=dict(color=C["body"], size=11),
            tickvals=[-1,-0.5,0,0.5,1],
            ticktext=["-1.0","-0.5","0.0","+0.5","+1.0"],
            thickness=14, len=0.8, bgcolor=C["bg"], bordercolor=C["border"]
        )
    ))
    fig_heat.update_layout(
        paper_bgcolor=C["bg"], plot_bgcolor=C["bg"],
        height=530, margin=dict(t=20,b=80,l=80,r=20),
        font=dict(color=C["body"], family="Inter, sans-serif"),
        hovermode="closest"
    )
    fig_heat.update_xaxes(tickangle=-40, tickfont=dict(size=12,color=C["body"]),
                          side="bottom", showgrid=False, linecolor=C["border"])
    fig_heat.update_yaxes(tickfont=dict(size=12,color=C["body"]),
                          autorange="reversed", showgrid=False, linecolor=C["border"])
    st.plotly_chart(fig_heat, use_container_width=True)

    flat = cvals[np.triu_indices_from(cvals, k=1)]
    s1,s2,s3,s4 = st.columns(4)
    s1.metric("Average Correlation", f"{flat.mean():.4f}", delta="Lower = more diversification")
    s2.metric("Max Correlation",     f"{flat.max():.4f}", delta="Most similar pair")
    s3.metric("Min Correlation",     f"{flat.min():.4f}", delta="Most different pair")
    s4.metric("Negative Pairs",      f"{(flat<0).sum()} / {len(flat)}", delta="Pairs with hedge benefit")

    with st.expander("📖 How to read this chart"):
        st.markdown("""
        - 🔴 **Red** = high positive correlation → stocks move together → less diversification
        - 🔵 **Blue** = low/negative correlation → stocks move independently → more diversification
        - **Diagonal** is always 1.0 (stock with itself)
        """)


# ════════════════════════════════════════════
#  TAB 5 — STOCK VALUATION
# ════════════════════════════════════════════
with tab5:
    st.markdown("<p style='font-size:14px; font-weight:700; color:#0f172a; margin-bottom:2px;'>📉 Individual Stock Valuation — CAPM vs Actual Return</p>", unsafe_allow_html=True)

    rows = []
    for t in available:
        b      = betas[t]
        actual = float(mean_ret[t])
        capm_r = rf_rate + b*(mkt_annual - rf_rate)
        ai     = actual - capm_r
        rows.append({"Ticker":t, "Label":short(t), "Beta":b,
                     "Actual":actual, "CAPM":capm_r, "Alpha":ai,
                     "Under": actual > capm_r,
                     "Weight": opt_w[available.index(t)]})

    vdf   = pd.DataFrame(rows).sort_values("Alpha", ascending=False)
    un_df = vdf[vdf["Under"]]
    ov_df = vdf[~vdf["Under"]]

    def val_card(row, is_under):
        color    = C["green"] if is_under else C["red"]
        brd      = "#bbf7d0" if is_under else "#fecaca"
        bg_badge = "#dcfce7" if is_under else "#fee2e2"
        return f"""
        <div style='background:#ffffff; border:1px solid {brd};
                    border-radius:10px; padding:13px 16px; margin-bottom:8px;
                    box-shadow:0 1px 3px rgba(0,0,0,.05);'>
            <div style='display:flex; justify-content:space-between;
                        align-items:center; margin-bottom:8px;'>
                <span style='color:{color}; font-weight:800; font-size:15px;'>
                    {row["Label"]}</span>
                <span style='background:{bg_badge}; color:{color}; padding:3px 10px;
                             border-radius:999px; font-size:12px; font-weight:700;'>
                    α = {row["Alpha"]:+.2%}</span>
            </div>
            <div style='display:grid; grid-template-columns:repeat(4,1fr); gap:6px;'>
                <div><div style='color:#64748b;font-size:10px;'>Beta</div>
                     <div style='color:#0f172a;font-weight:700;font-size:13px;'>{row["Beta"]:.4f}</div></div>
                <div><div style='color:#64748b;font-size:10px;'>Actual</div>
                     <div style='color:{color};font-weight:700;font-size:13px;'>{row["Actual"]:.2%}</div></div>
                <div><div style='color:#64748b;font-size:10px;'>CAPM Req.</div>
                     <div style='color:{C["amber"]};font-weight:700;font-size:13px;'>{row["CAPM"]:.2%}</div></div>
                <div><div style='color:#64748b;font-size:10px;'>Port. Wt</div>
                     <div style='color:{C["blue"]};font-weight:700;font-size:13px;'>{row["Weight"]:.2%}</div></div>
            </div>
        </div>"""

    hc1, hc2 = st.columns(2, gap="large")

    with hc1:
        st.markdown(f"""
        <div style='background:#f0fdf4; border:1px solid #bbf7d0; border-radius:12px;
                    padding:12px 16px; margin-bottom:12px;'>
            <span style='color:{C["green"]}; font-weight:800; font-size:15px;'>
                ✅ Undervalued — {len(un_df)} stock{"s" if len(un_df)!=1 else ""}</span>
            <div style='color:#15803d; font-size:12px; margin-top:3px;'>
                Actual return &gt; CAPM → Positive Alpha → Above SML</div>
        </div>""", unsafe_allow_html=True)
        for _, row in un_df.iterrows():
            st.markdown(val_card(row, True), unsafe_allow_html=True)

    with hc2:
        st.markdown(f"""
        <div style='background:#fef2f2; border:1px solid #fecaca; border-radius:12px;
                    padding:12px 16px; margin-bottom:12px;'>
            <span style='color:{C["red"]}; font-weight:800; font-size:15px;'>
                ❌ Overvalued — {len(ov_df)} stock{"s" if len(ov_df)!=1 else ""}</span>
            <div style='color:#b91c1c; font-size:12px; margin-top:3px;'>
                Actual return &lt; CAPM → Negative Alpha → Below SML</div>
        </div>""", unsafe_allow_html=True)
        for _, row in ov_df.iterrows():
            st.markdown(val_card(row, False), unsafe_allow_html=True)

    st.markdown("<p style='font-size:14px; font-weight:700; color:#0f172a; margin:16px 0 6px;'>📊 Jensen's Alpha — All Stocks</p>", unsafe_allow_html=True)
    vs = vdf.sort_values("Alpha")
    fig_a = go.Figure(go.Bar(
        x=vs["Label"], y=vs["Alpha"]*100,
        marker=dict(color=[C["green"] if a>0 else C["red"] for a in vs["Alpha"]],
                    opacity=0.85, line=dict(color="rgba(255,255,255,0.6)", width=1)),
        text=[f"{a:+.2f}%" for a in vs["Alpha"]*100],
        textposition="outside", textfont=dict(color=C["head"], size=11),
        hovertemplate="<b>%{x}</b><br>Alpha: %{y:.2f}%<extra></extra>"
    ))
    fig_a.add_hline(y=0, line_color=C["muted"], line_width=1.5, line_dash="dot")
    apply_white_theme(fig_a, height=310, margin=dict(t=20,b=20,l=10,r=10))
    style_axes(fig_a,
               xkw=dict(tickfont=dict(size=12, color=C["body"])),
               ykw=dict(title="Jensen's Alpha (%)", ticksuffix="%"))
    fig_a.update_layout(showlegend=False, hovermode="closest")
    st.plotly_chart(fig_a, use_container_width=True)


# ════════════════════════════════════════════
#  TAB 6 — CAL SIMULATION
# ════════════════════════════════════════════
with tab6:
    st.markdown("<p style='font-size:14px; font-weight:700; color:#0f172a; margin-bottom:2px;'>📋 Capital Allocation Line — 101 Portfolio Combinations</p>", unsafe_allow_html=True)
    st.caption("Shifting 1% at a time: 100% Risk-Free → 100% Optimal Risky Portfolio")

    w_arr = np.linspace(0, 1, 101)
    r_arr = rf_rate + w_arr*(opt_ret - rf_rate)
    v_arr = w_arr * opt_vol

    fig_lines = make_subplots(
        rows=1, cols=2,
        subplot_titles=["Expected Return vs Risky Weight",
                        "Volatility vs Risky Weight"]
    )
    fig_lines.add_trace(go.Scatter(
        x=w_arr*100, y=r_arr*100, mode="lines",
        line=dict(color=C["green"], width=2.5), name="Return",
        hovertemplate="Risky Wt: %{x:.0f}%<br>Return: %{y:.2f}%<extra></extra>"
    ), 1, 1)
    fig_lines.add_trace(go.Scatter(
        x=w_arr*100, y=v_arr*100, mode="lines",
        line=dict(color=C["amber"], width=2.5), name="Volatility",
        hovertemplate="Risky Wt: %{x:.0f}%<br>Vol: %{y:.2f}%<extra></extra>"
    ), 1, 2)

    fig_lines.update_layout(
        paper_bgcolor=C["bg"], plot_bgcolor=C["bg"],
        height=270, margin=dict(t=40,b=20,l=10,r=10),
        showlegend=False, font=dict(color=C["body"], family="Inter, sans-serif")
    )
    fig_lines.update_xaxes(gridcolor=C["border"], zeroline=False, ticksuffix="%",
                           tickfont=dict(size=10, color=C["muted"]),
                           title_text="Risky Weight (%)",
                           title_font=dict(color=C["body"], size=11))
    fig_lines.update_yaxes(gridcolor=C["border"], zeroline=False,
                           tickfont=dict(size=10, color=C["muted"]))
    for ann in fig_lines.layout.annotations:
        ann.font.color = C["head"]; ann.font.size = 12

    st.plotly_chart(fig_lines, use_container_width=True)

    cal_rows = []
    for i in range(101):
        wrf=i/100; wr=1-wrf
        cr = wrf*rf_rate + wr*opt_ret
        cv = wr*opt_vol
        cb = wr*port_beta
        cc = rf_rate + cb*(mkt_annual - rf_rate)
        cs = (cr-rf_rate)/cv if cv > 0 else 0.0
        cal_rows.append({
            "Risk-Free Wt":     f"{wrf:.0%}",
            "Risky Wt":         f"{wr:.0%}",
            "Exp. Return":      f"{cr:.2%}",
            "Volatility":       f"{cv:.2%}",
            "Beta":             f"{cb:.4f}",
            "CAPM Req. Return": f"{cc:.2%}",
            "Sharpe Ratio":     f"{cs:.4f}",
        })
    st.dataframe(pd.DataFrame(cal_rows),
                 use_container_width=True, hide_index=True, height=420)


# ════════════════════════════════════════════
#  TAB 7 — EXCEL EXPORT (only if checkbox ticked)
# ════════════════════════════════════════════
if want_excel and tab7 is not None:
    with tab7:
        st.caption("Generates a formula-driven 3-sheet Excel workbook")
        with st.spinner("Building Excel..."):
            xl = build_excel(
                stock_data=stock_data, market_data=market_data,
                returns=returns, mkt_ret=mkt_ret,
                mean_ret=mean_ret, cov_mat=cov_mat, corr_mat=corr_mat,
                mkt_annual=mkt_annual, available=available,
                market_choice=market_choice, rf_rate=rf_rate,
                opt_w=opt_w, opt_ret=opt_ret, opt_vol=opt_vol, opt_sh=opt_sh,
                port_beta=port_beta, capm_ret=capm_ret, alpha_val=alpha,
                betas=betas, gmvp_ret=gmvp_ret, gmvp_vol=gmvp_vol,
                start_date=start_date, end_date=end_date
            )
        if xl:
            st.download_button(
                label="📥  Download Excel Report",
                data=xl,
                file_name=f"Portfolio_Analysis_{start_date}_{end_date}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
        else:
            st.error("Excel generation failed. Ensure `openpyxl` is installed.")


# ─────────────────────────────────────────────
#  FOOTER
# ─────────────────────────────────────────────
st.markdown("---")
st.markdown(f"""
<div style='text-align:center; color:{C["muted"]}; font-size:12px; padding:8px 0;'>
    📊 Portfolio Optimizer &nbsp;·&nbsp; Data: Yahoo Finance &nbsp;·&nbsp;
    {start_date} → {end_date} &nbsp;·&nbsp; Rf: {rf_rate:.1%} &nbsp;·&nbsp;
    Market: {market_choice} &nbsp;·&nbsp; {len(available)} assets via SLSQP
    <br><span style='color:#94a3b8; font-size:11px; margin-top:4px; display:block;'>
        ⚠️ Educational purposes only. Past performance does not guarantee future results.
    </span>
</div>
""", unsafe_allow_html=True)
